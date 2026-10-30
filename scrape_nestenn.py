"""
Récupère l'e-mail général (et SIRET / SIREN) de chaque agence Nestenn
en lisant la page /cgu de son site (ex. https://immobilier-agen.nestenn.com/cgu).

- Entrée / sortie : agences_nestenn.xlsx (feuille "Agences")
- Colonnes utilisées : "Site internet", "Email général", "SIRET", "SIREN"
- Les valeurs déjà remplies ne sont jamais écrasées.
"""
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

FICHIER = os.environ.get("FICHIER", "agences_nestenn.xlsx")
FEUILLE = os.environ.get("FEUILLE", "Agences")
WORKERS = int(os.environ.get("WORKERS", "5"))
PAUSE = float(os.environ.get("PAUSE", "0.5"))
TIMEOUT = 20

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "fr-FR,fr;q=0.9",
}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Adresses génériques du réseau, à ne pas prendre pour l'e-mail de l'agence
EMAILS_IGNORES = {"web@nestenn.com", "contact@nestenn.com"}


def charger_page(session, url):
    for tentative in range(3):
        try:
            r = session.get(url, headers=HEADERS, timeout=TIMEOUT)
            if r.status_code == 200:
                return r.text
            if r.status_code in (404, 410):
                return None
        except requests.RequestException:
            pass
        time.sleep(2 * (tentative + 1))
    return None


def valeur_apres(lignes, etiquette):
    """Retourne la ligne qui suit une étiquette du type 'Mail :'."""
    for i, ligne in enumerate(lignes[:-1]):
        if ligne.lower().rstrip(" :") == etiquette.lower():
            return lignes[i + 1]
    return None


def extraire(html):
    soupe = BeautifulSoup(html, "html.parser")
    for t in soupe(["script", "style"]):
        t.decompose()
    lignes = [l for l in soupe.get_text("\n", strip=True).split("\n") if l]

    # 1) E-mail : bloc "Siège social ... Mail : xxx@nestenn.com"
    email = None
    brut = valeur_apres(lignes, "Mail")
    if brut:
        m = EMAIL_RE.search(brut)
        if m:
            email = m.group(0).lower()
    # 2) Repli : première adresse qui n'est pas une adresse générique du réseau
    if not email:
        for m in EMAIL_RE.finditer("\n".join(lignes)):
            candidat = m.group(0).lower().rstrip(".")
            if candidat not in EMAILS_IGNORES:
                email = candidat
                break

    # SIRET / SIREN : le champ "RCS" contient le SIRET (14 chiffres)
    siret = siren = None
    rcs = valeur_apres(lignes, "RCS")
    if rcs:
        chiffres = re.sub(r"\D", "", rcs)
        if len(chiffres) == 14:
            siret, siren = chiffres, chiffres[:9]
        elif len(chiffres) == 9:
            siren = chiffres
    return email, siret, siren


def traiter(session, index, site):
    url = site.rstrip("/") + "/cgu"
    html = charger_page(session, url)
    time.sleep(PAUSE)
    if not html:
        return index, url, None, None, None
    return (index, url, *extraire(html))


def main():
    wb = load_workbook(FICHIER)
    ws = wb[FEUILLE]
    cols = {c.value: c.column for c in ws[1] if c.value}
    for nom in ("Site internet", "Email général"):
        if nom not in cols:
            raise SystemExit(f"Colonne manquante : {nom}")
    c_site, c_mail = cols["Site internet"], cols["Email général"]
    c_siret, c_siren = cols.get("SIRET"), cols.get("SIREN")

    a_faire = []
    for row in range(2, ws.max_row + 1):
        site = ws.cell(row, c_site).value
        mail = ws.cell(row, c_mail).value
        if site and not mail and "nestenn.com" in str(site):
            a_faire.append((row, str(site).strip()))
    print(f"{len(a_faire)} agences à traiter")

    session = requests.Session()
    trouves = 0
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = [pool.submit(traiter, session, r, s) for r, s in a_faire]
        for n, f in enumerate(as_completed(futures), 1):
            row, url, email, siret, siren = f.result()
            if email:
                ws.cell(row, c_mail).value = email
                trouves += 1
            if siret and c_siret and not ws.cell(row, c_siret).value:
                ws.cell(row, c_siret).value = siret
            if siren and c_siren and not ws.cell(row, c_siren).value:
                ws.cell(row, c_siren).value = siren
            print(f"[{n}/{len(a_faire)}] {url} -> {email or 'rien trouvé'}")
            if n % 25 == 0:
                wb.save(FICHIER)  # sauvegarde intermédiaire

    wb.save(FICHIER)
    print(f"Terminé : {trouves} e-mails trouvés sur {len(a_faire)} agences")


if __name__ == "__main__":
    main()
