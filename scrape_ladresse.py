"""
Récupère, pour chaque agence L'Adresse, tous les membres de l'équipe
(nom, prénom, fonction) depuis la page agence sur ladresse.com,
ainsi que le représentant légal (gérant) et le SIREN (via le numéro RCS).

- Entrée / sortie : ladresse_gerants.xlsx
    * feuille "Agences"  : ID agence, Nom agence, Site internet, SIREN,
                           Nom du gérant, Prénom du gérant, E-mail du gérant
    * feuille "Contacts" : créée / complétée par le script, une ligne par membre
- Les agences déjà présentes dans "Contacts" sont ignorées (le script est relançable).
- Les e-mails ne sont pas publiés sur ces pages : la colonne reste vide.
"""
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

FICHIER = os.environ.get("FICHIER", "ladresse_gerants.xlsx")
FEUILLE = os.environ.get("FEUILLE", "Agences")
FEUILLE_CONTACTS = "Contacts"
WORKERS = int(os.environ.get("WORKERS", "3"))
PAUSE = float(os.environ.get("PAUSE", "0.5"))
TIMEOUT = 25

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "fr-FR,fr;q=0.9",
}

COLONNES_CONTACTS = [
    "ID agence", "Nom agence", "SIREN", "Nom", "Prénom", "Fonction",
    "Email professionnel", "Source",
]
MOTS_GERANT = ("gérant", "gerant", "dirigeant", "directeur", "directrice")


def charger_page(session, url):
    for tentative in range(3):
        try:
            r = session.get(url, headers=HEADERS, timeout=TIMEOUT)
            if r.status_code == 200:
                return r.text
            if r.status_code in (404, 410):
                return None
        except requests.RequestException:
            pass
        time.sleep(2 * (tentative + 1))
    return None


def separer_nom_prenom(texte):
    """'Raphaëlle DUMAND' ou 'CHARBONNEL Guillaume' -> (nom, prénom)."""
    jetons = texte.split()
    if not jetons:
        return "", ""
    maj = [j.isupper() for j in jetons]
    if not any(maj) or all(maj):          # pas de repère : 1er mot = prénom
        prenom, nom = jetons[:1], jetons[1:]
    elif maj[0]:                          # NOM Prénom
        i = next(k for k, m in enumerate(maj) if not m)
        nom, prenom = jetons[:i], jetons[i:]
    else:                                 # Prénom NOM
        i = next(k for k, m in enumerate(maj) if m)
        prenom, nom = jetons[:i], jetons[i:]
    return " ".join(nom).title(), " ".join(prenom).title()


def valeur_apres(lignes, debut_etiquette):
    for i, ligne in enumerate(lignes[:-1]):
        if ligne.lower().rstrip(" :").startswith(debut_etiquette):
            return lignes[i + 1]
    return None


def extraire(html):
    soupe = BeautifulSoup(html, "html.parser")

    # Équipe : <div class="agent"><div class="agent-nom">..</div><div class="agent-fonction">..</div>
    membres, vus = [], set()
    for bloc in soupe.select("div.agent"):
        n = bloc.select_one(".agent-nom")
        f = bloc.select_one(".agent-fonction")
        nom_complet = n.get_text(" ", strip=True) if n else ""
        fonction = f.get_text(" ", strip=True) if f else ""
        if not nom_complet or (nom_complet, fonction) in vus:
            continue
        vus.add((nom_complet, fonction))
        nom, prenom = separer_nom_prenom(nom_complet)
        membres.append({"nom": nom, "prenom": prenom, "fonction": fonction})

    # Mentions légales du pied de page : représentant légal et numéro RCS
    for t in soupe(["script", "style"]):
        t.decompose()
    lignes = [l for l in soupe.get_text("\n", strip=True).split("\n") if l]
    representant = valeur_apres(lignes, "representant légal") or valeur_apres(
        lignes, "représentant légal"
    )
    siren = None
    rcs = valeur_apres(lignes, "numéro rcs")
    if rcs:
        chiffres = re.sub(r"\D", "", rcs)
        if len(chiffres) >= 9:
            siren = chiffres[-9:]

    gerant = None
    if representant:
        nom, prenom = separer_nom_prenom(representant)
        gerant = {"nom": nom, "prenom": prenom}
    else:
        for m in membres:
            if any(mot in m["fonction"].lower() for mot in MOTS_GERANT):
                gerant = {"nom": m["nom"], "prenom": m["prenom"]}
                break
    return membres, gerant, siren


def traiter(session, ligne, url):
    html = charger_page(session, url)
    time.sleep(PAUSE)
    if not html:
        return ligne, None
    return ligne, extraire(html)


def preparer_feuille_contacts(wb):
    if FEUILLE_CONTACTS in wb.sheetnames:
        return wb[FEUILLE_CONTACTS]
    ws = wb.create_sheet(FEUILLE_CONTACTS)
    ws.append(COLONNES_CONTACTS)
    return ws


def main():
    if not os.path.exists(FICHIER):
        raise SystemExit(
            f"Fichier introuvable : {FICHIER}. Ajoutez-le à la racine du dépôt "
            f"(fichiers présents : {sorted(os.listdir('.'))})"
        )
    wb = load_workbook(FICHIER)
    ws = wb[FEUILLE]
    cols = {c.value: c.column for c in ws[1] if c.value}
    for nom in ("ID agence", "Nom agence", "Site internet"):
        if nom not in cols:
            raise SystemExit(f"Colonne manquante : {nom}")
    c_id, c_nom, c_site = cols["ID agence"], cols["Nom agence"], cols["Site internet"]
    c_siren = cols.get("SIREN")
    c_gnom, c_gpre = cols.get("Nom du gérant"), cols.get("Prénom du gérant")

    wc = preparer_feuille_contacts(wb)
    deja = {wc.cell(r, 1).value for r in range(2, wc.max_row + 1)}

    a_faire = []
    for r in range(2, ws.max_row + 1):
        ident, site = ws.cell(r, c_id).value, ws.cell(r, c_site).value
        if site and ident not in deja:
            a_faire.append((r, str(site).strip()))
    print(f"{len(a_faire)} agences à traiter")

    session = requests.Session()
    total, echecs = 0, 0
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = [pool.submit(traiter, session, r, u) for r, u in a_faire]
        for n, f in enumerate(as_completed(futures), 1):
            r, resultat = f.result()
            ident, nom_ag = ws.cell(r, c_id).value, ws.cell(r, c_nom).value
            url = ws.cell(r, c_site).value
            if resultat is None:
                echecs += 1
                print(f"[{n}/{len(a_faire)}] {nom_ag} : page non récupérée")
                continue
            membres, gerant, siren = resultat
            if siren and c_siren and not ws.cell(r, c_siren).value:
                ws.cell(r, c_siren).value = siren
            if gerant and c_gnom and not ws.cell(r, c_gnom).value:
                ws.cell(r, c_gnom).value = gerant["nom"]
                if c_gpre:
                    ws.cell(r, c_gpre).value = gerant["prenom"]
            for m in membres:
                wc.append([ident, nom_ag, siren, m["nom"], m["prenom"],
                           m["fonction"], None, url])
            if not membres:  # marque l'agence comme traitée (aucune équipe affichée)
                wc.append([ident, nom_ag, siren, None, None, "Aucun membre affiché", None, url])
            total += len(membres)
            print(f"[{n}/{len(a_faire)}] {nom_ag} : {len(membres)} membres")
            if n % 25 == 0:
                wb.save(FICHIER)

    wb.save(FICHIER)
    print(f"Terminé : {total} contacts, {echecs} pages non récupérées")


if __name__ == "__main__":
    main()
